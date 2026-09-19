from pathlib import Path
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import build_report as b


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "wise-revolut-monarch-deep-dive.docx"
SHOTS = ROOT / "screenshots"


def title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(74)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("WOLLIE  /  COMPETITIVE DEEP DIVE")
    b.set_run(r, 10, color=b.ORANGE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    b.set_run(r, 29, color=b.NAVY, bold=True)
    p.paragraph_format.line_spacing = 1.02
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(34)
        r = p.add_run(subtitle)
        b.set_run(r, 14, color=b.MUTED)

    t = doc.add_table(rows=1, cols=1)
    b.set_table_geometry(t, [b.PAGE_WIDTH_DXA])
    c = t.cell(0, 0)
    b.set_cell_shading(c, b.CREAM)
    b.set_cell_border(c, top={"val": "single", "sz": "12", "color": b.ORANGE}, bottom={"val": "single", "sz": "12", "color": b.ORANGE})
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Wise  /  Revolut  /  Monarch\nResearch current to 1 September 2026")
    b.set_run(r, 11, color=b.NAVY, bold=True)
    b.set_paragraph_spacing(p, after=0)

    b.add_text(doc, "Official product documentation and annual reports are the primary evidence. Public user discussions are labeled as anecdotal signals.", size=10.5, color=b.MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=35, after=0)
    doc.add_page_break()


def page_intro(doc):
    b.add_heading(doc, "The answer in one page", 1, before=0)
    b.add_text(doc, "These products are not three versions of the same app. Wise is an international money account, Revolut is a digital banking super-app, and Monarch is a cross-bank planning layer. Wollie should combine their strongest interaction patterns without combining their complexity.", size=12, color=b.NAVY, after=10)
    b.add_callout(doc, "Recommended position", "A calm, couples-first money planning layer across joint, separate, or mixed bank accounts. It coordinates the household without pretending a virtual plan moved real money.", b.PALE_ORANGE)

    cards = doc.add_table(rows=1, cols=3)
    b.set_table_geometry(cards, [3120, 3120, 3120])
    data = [
        ("WISE", "Clarity", "Borrow its balance, transaction, status, and fee hierarchy."),
        ("REVOLUT", "Immediacy", "Borrow rule previews, instant feedback, and real joint boundaries."),
        ("MONARCH", "Planning", "Borrow household aggregation, review queues, rules, goals, and reports."),
    ]
    for cell, (kicker, heading, body) in zip(cards.rows[0].cells, data):
        b.set_cell_shading(cell, b.PALE_BLUE)
        b.set_cell_border(cell, top={"val": "single", "sz": "6", "color": b.BLUE}, bottom={"val": "single", "sz": "6", "color": b.BLUE})
        p = cell.paragraphs[0]
        r = p.add_run(kicker + "\n")
        b.set_run(r, 8.5, color=b.ORANGE, bold=True)
        r = p.add_run(heading + "\n")
        b.set_run(r, 13, color=b.NAVY, bold=True)
        r = p.add_run(body)
        b.set_run(r, 9.2, color=b.INK)
        b.set_paragraph_spacing(p, after=0)

    b.add_heading(doc, "The non-negotiable model", 2)
    t = doc.add_table(rows=3, cols=2)
    b.set_table_geometry(t, [4680, 4680])
    for c, text in zip(t.rows[0].cells, ["BANK REALITY", "WOLLIE PLAN"]):
        c.text = text
    for row, values in zip(t.rows[1:], [
        ("Balances, transactions, transfers, pending and cleared states.", "Virtual assignments to spending, tax, pension, savings, and goals."),
        ("Answers: what actually moved?", "Answers: what is this month's money meant to do?"),
    ]):
        for c, value in zip(row.cells, values):
            c.text = value
    b.style_table(t)
    b.add_text(doc, "A pension target is planned until a matching outgoing contribution is detected or confirmed. It must not appear as everyday spending.", size=10.5, color=b.MUTED, italic=True, after=0)
    doc.add_page_break()


def page_position(doc):
    b.add_heading(doc, "Category map", 1, before=0)
    b.add_text(doc, "The correct benchmark depends on the job being evaluated.", color=b.MUTED, after=8)
    t = doc.add_table(rows=1, cols=5)
    b.set_table_geometry(t, [1250, 1700, 1950, 2020, 2440])
    for c, text in zip(t.rows[0].cells, ["PRODUCT", "PRIMARY JOB", "HOUSEHOLD MODEL", "BEST AT", "MAIN GAP"]):
        c.text = text
    rows = [
        ("Wise", "Hold, convert, send, spend", "Owner-led shared spending", "Clear ledger and international money movement", "Not a household planning system"),
        ("Revolut", "Daily digital banking", "Two-person joint account plus personal accounts", "Actual-money automation and control", "Feature and analytics fragmentation"),
        ("Monarch", "Cross-bank planning", "Shared household subscription", "Review, budgets, goals, reports", "Sync trust and no real privacy"),
        ("Wollie", "Couples money coordination", "Joint, separate, or mixed banks", "Fair contributions plus selective sharing", "Must prove data and plan accuracy"),
    ]
    for values in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, values):
            c.text = value
    b.style_table(t)

    b.add_heading(doc, "What each product optimizes", 2)
    for text in [
        "Wise optimizes a transfer or balance event. Its success metric is confidence that the money arrived at a fair cost.",
        "Revolut optimizes action and engagement. Its success metric is how much of a person's financial life happens inside Revolut.",
        "Monarch optimizes understanding and planning. Its success metric is whether connected data becomes a coherent monthly and long-term picture.",
        "Wollie should optimize household agreement. Its success metric is whether both partners understand what is available, what is reserved, and what needs attention.",
    ]:
        b.add_bullet(doc, text)
    b.add_callout(doc, "Design consequence", "Do not benchmark Wollie's Overview against Revolut's entire home screen. Benchmark Accounts against Wise, automation against Revolut, and planning/review against Monarch.", b.PALE_GREEN)
    doc.add_page_break()


def page_product(doc, name, subtitle, image, image_caption, bullets, strengths, weaknesses):
    b.add_heading(doc, name, 1, before=0)
    b.add_text(doc, subtitle, size=12, color=b.NAVY, after=8)
    b.add_picture_with_caption(doc, image, image_caption, width=5.95)
    b.add_heading(doc, "How the product thinks", 2)
    for item in bullets:
        b.add_bullet(doc, item)

    t = doc.add_table(rows=1, cols=2)
    b.set_table_geometry(t, [4680, 4680])
    t.rows[0].cells[0].text = "STRENGTHS"
    t.rows[0].cells[1].text = "WEAKNESSES"
    max_rows = max(len(strengths), len(weaknesses))
    for i in range(max_rows):
        cells = t.add_row().cells
        cells[0].text = strengths[i] if i < len(strengths) else ""
        cells[1].text = weaknesses[i] if i < len(weaknesses) else ""
    b.style_table(t)
    doc.add_page_break()


def wise_second(doc):
    b.add_heading(doc, "Wise: shared money and planning logic", 1, before=0)
    b.add_picture_with_caption(doc, SHOTS / "wise-spend-with-others-ui.png", "Wise's public Spend with others experience emphasizes a shared spending event, not a household budget.", width=2.55)
    b.add_heading(doc, "Spend with others is controlled access", 2)
    b.add_text(doc, "One owner legally owns and controls all money, including member contributions. Members may add money, see transactions, and spend, but cannot send, withdraw, convert, or manage the group. The current EU group limit is two people. This is useful for controlled shared expenses but is not an equal joint account.")
    b.add_heading(doc, "Jars are real-money containers", 2)
    b.add_text(doc, "Money in a Jar is separated from the main spendable balance and cannot be used by the debit card or direct debit. This makes the number concrete and prevents accidental spending. It also requires the money to live inside Wise.")
    b.add_heading(doc, "Implication for Wollie", 2)
    b.add_callout(doc, "Borrow", "A simple account screen, explicit statuses, and an unmistakable difference between available money and set-aside money.", b.PALE_GREEN)
    b.add_callout(doc, "Avoid", "Presenting an owner/member permission model as equal household collaboration, or requiring couples to move shared money into Wollie.", b.PALE_ORANGE)
    doc.add_page_break()


def revolut_second(doc):
    b.add_heading(doc, "Revolut: joint accounts, automation, and analytics", 1, before=0)
    b.add_picture_with_caption(doc, SHOTS / "revolut-joint-phone.png", "Revolut uses immersive imagery and a high-contrast spending card to make analytics feel immediate and personal.", width=2.7)
    b.add_heading(doc, "The joint account is a separate real account", 2)
    b.add_text(doc, "Both holders can add, spend, withdraw, and manage money. Each must already have a verified personal Revolut account, be under the same country/entity, be contacts, and accept the invitation. Joint spending uses cards linked to that account. The clean ledger boundary is useful, but the setup can feel like creating a third account.")
    b.add_heading(doc, "Income Sorter moves actual money", 2)
    b.add_text(doc, "When recognized income arrives, users can distribute fixed amounts to destinations such as a joint account, Pockets, Savings, Investments, or another person. This is a strong automation pattern because the trigger, action, and result are visible.")
    b.add_heading(doc, "Analytics has a household gap", 2)
    b.add_text(doc, "Revolut supports spending, income, cash flow, total wealth, categories, merchants, countries, currencies, and cards. Current official help material states that joint accounts are excluded from the spending section, so personal analytics does not automatically become one household truth.")
    b.add_callout(doc, "Wollie opportunity", "Use household-wide proportional allocation across multiple incomes. Preview the calculation, then keep virtual assignment separate from any real transfer.", b.PALE_GREEN)
    doc.add_page_break()


def monarch_first(doc):
    b.add_heading(doc, "Monarch: the closest planning benchmark", 1, before=0)
    b.add_text(doc, "Monarch connects external accounts and turns their data into a household workspace. It is the closest direct reference for Wollie's transactions, plans, goals, recurring expenses, and reports.", size=12, color=b.NAVY, after=8)
    b.add_picture_with_caption(doc, SHOTS / "monarch-couples-ui-1.png", "Monarch combines account aggregation, ownership labels, and net worth in one household dashboard.", width=5.35)
    b.add_heading(doc, "Why it matters", 2)
    for text in [
        "Separate logins and one subscription support a real household workflow.",
        "Connected data is transformed through categories, tags, rules, review assignments, budgets, goals, and reports.",
        "The product gives value before a budget is complete: cash flow and categorized spending still work.",
        "Its weaknesses expose the hardest parts of Wollie's problem: sync trust, privacy, and explainable plan maths.",
    ]:
        b.add_bullet(doc, text)
    doc.add_page_break()


def monarch_shared(doc):
    b.add_heading(doc, "Monarch: collaboration without privacy", 1, before=0)
    b.add_picture_with_caption(doc, SHOTS / "monarch-transaction-owner.png", "Shared Views assigns an owner to a transaction. Ownership changes filtering and attribution, not access permission.", width=5.95)
    b.add_heading(doc, "Shared Views", 2)
    b.add_text(doc, "Accounts and transactions can be Shared or Individual and filtered by owner across Accounts, Transactions, Reports, and Cash Flow. Ownership is inherited from an account and may be changed with rules.")
    b.add_callout(doc, "Critical limitation", "Monarch's help centre says all household members still have full visibility. Shared and Individual are labels and filters, not privacy controls. The household also has one shared budget.", b.PALE_ORANGE)
    b.add_heading(doc, "Lifecycle friction", 2)
    b.add_text(doc, "A person with an existing Monarch account cannot simply join another household with that account. Official guidance points to using another email/new account or deleting the old account. Removing a member can also remove accounts they added.")
    b.add_heading(doc, "Wollie model", 2)
    b.add_text(doc, "Keep legal account ownership, transaction attribution, and visibility as separate fields. Visibility must be enforceable: Private, Shared details, or Household total only.", color=b.NAVY, bold=True, after=0)
    doc.add_page_break()


def monarch_review(doc):
    b.add_heading(doc, "Monarch: review, rules, and recurring", 1, before=0)
    b.add_picture_with_caption(doc, SHOTS / "monarch-couples-ui-2.png", "A household member can own the review task for uncertain transactions. This creates a practical collaboration loop.", width=5.8)
    b.add_heading(doc, "Transactions and rules", 2)
    b.add_text(doc, "Each transaction has one category and may have multiple tags. Rules can match merchant, amount, category, or account, then rename, categorize, tag, hide, mark for review, or link to a goal. Rule order matters.")
    b.add_heading(doc, "Needs Review", 2)
    b.add_text(doc, "Transactions can be assigned to self, partner, or another household member. This is one of Monarch's best patterns because it makes imperfect automation collaborative instead of invisible.")
    b.add_heading(doc, "Recurring detection", 2)
    b.add_text(doc, "Monarch detects recurring merchants, transfers, and paychecks and shows expected dates and status. The model is valuable for forecasting but probabilistic. Wollie should expose confidence, correction, and a link back to supporting transactions.")
    b.add_callout(doc, "Simplify for Wollie", "Ask one review question at a time. Offer Apply once or Create rule. Always provide undo and Why did Wollie do this?", b.PALE_GREEN)
    doc.add_page_break()


def monarch_budget(doc):
    b.add_heading(doc, "Monarch: budget and goals logic", 1, before=0)
    b.add_picture_with_caption(doc, SHOTS / "monarch-budget-overview.png", "Monarch's budget summary leads with Left to budget, a useful but calculation-heavy concept.", width=5.3)
    b.add_heading(doc, "Category and Flex budgets", 2)
    b.add_text(doc, "The budget is monthly cash-flow based rather than a mirror of account balances. Expected income, planned expenses, and savings determine the monthly picture. Flex budgeting groups fixed, non-monthly, and flexible spending. A category can also have no target and still track spending, which is useful for unpredictable Health expenses.")
    b.add_heading(doc, "Rollover", 2)
    b.add_text(doc, "Rollover carries the previous remaining amount into the new month. This is powerful for irregular expenses, but it adds another state users must reconcile with the bank balance.")
    b.add_heading(doc, "Goals", 2)
    b.add_text(doc, "Goals can link planned contributions and account balances. Public product changes through 2026 show how difficult the model is to get right. The main risk is blending a planned contribution, an actual transfer, and a changing asset balance.")
    b.add_callout(doc, "Wollie rule", "Every reserve or goal shows Planned this month, Actually contributed, and Current connected balance as separate facts.", b.PALE_GREEN)
    doc.add_page_break()


def journeys(doc):
    b.add_heading(doc, "Five journeys Wollie must make clearer", 1, before=0)
    rows = [
        ("1. Salary arrives", "Show recognized income, proportional allocation preview, and new plan amounts. Do not imply money moved."),
        ("2. Grocery appears", "Categorize to Groceries, map to Food, and reduce the amount left. Keep the transaction row compact."),
        ("3. Pension transfer", "Count as an actual contribution, not everyday spending. Preserve the original bank transaction."),
        ("4. Partner joins", "Explain sign in or create account, permissions, visibility, and existing-account rules before connection."),
        ("5. Data is uncertain", "Send only the uncertain field to Needs review, with a rule preview and undo."),
    ]
    t = doc.add_table(rows=1, cols=2)
    b.set_table_geometry(t, [2500, 6860])
    t.rows[0].cells[0].text = "JOURNEY"
    t.rows[0].cells[1].text = "IDEAL WOLLIE BEHAVIOUR"
    for label, body in rows:
        cells = t.add_row().cells
        cells[0].text = label
        cells[1].text = body
    b.style_table(t)
    b.add_heading(doc, "Default income allocation", 2)
    b.add_text(doc, "Allocation to item = received income × item remaining gap ÷ total remaining gaps", size=12, color=b.NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=8)
    b.add_text(doc, "Example: if €10,000 arrives against €17,000 of remaining targets, each target receives 58.8% of its remaining gap. Later income continues filling the gaps. No hidden priority order is created.")
    b.add_heading(doc, "Transfer and card-payment rule", 2)
    b.add_text(doc, "Card purchases count as spending. Paying the card bill does not count again. Transfers to pension, savings, or investments count as contributions when matched, not as everyday spending.")
    b.add_callout(doc, "Traceability", "Every summary amount must open to the transactions or calculations that compose it. Trust is a product feature.", b.PALE_ORANGE)
    doc.add_page_break()


def information_architecture(doc):
    b.add_heading(doc, "Recommended Wollie information architecture", 1, before=0)
    b.add_text(doc, "Keep five primary destinations. Put household, pricing, and settings in secondary navigation.", size=12, color=b.NAVY, after=8)
    t = doc.add_table(rows=1, cols=3)
    b.set_table_geometry(t, [1700, 3000, 4660])
    for c, text in zip(t.rows[0].cells, ["DESTINATION", "PRIMARY QUESTION", "CORE CONTENT"]):
        c.text = text
    rows = [
        ("Overview", "Where do we stand?", "Income received, everyday spending left, saved/invested, needs review, upcoming exceptions"),
        ("Transactions", "What happened?", "Merchant, amount, category, account, status, attribution, category totals"),
        ("Money plan", "What should this month's money do?", "Tabs: Spending this month; Savings & reserves"),
        ("Recurring", "What is expected next?", "Expected date/amount, status, confidence, supporting transactions"),
        ("Accounts", "Where does the money live?", "Institution, owner, visibility, balance, sync time, connection health"),
    ]
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            c.text = value
    b.style_table(t)
    b.add_heading(doc, "Overview limit", 2)
    b.add_text(doc, "Use no more than four top-level facts. Suggested default: Income received, Everyday spending left, Saved or invested, and Needs review. Assigned or set aside totals belong inside Money plan unless they answer an immediate decision.")
    b.add_heading(doc, "Transaction list rule", 2)
    b.add_text(doc, "Do not repeat a full envelope sentence under every transaction. Show a category chip in the list. Put the detailed plan effect in transaction detail.")
    b.add_heading(doc, "Bills versus Money plan", 2)
    b.add_text(doc, "Recurring is a forecast of expected transactions. Money plan is the allowance or reserve. Rent may exist in both without duplication because one answers when payment is expected and the other answers how much is available.")
    doc.add_page_break()


def visual_principles(doc):
    b.add_heading(doc, "Visual and interaction principles", 1, before=0)
    b.add_heading(doc, "One colour, one meaning", 2)
    t = doc.add_table(rows=1, cols=3)
    b.set_table_geometry(t, [1500, 2500, 5360])
    for c, text in zip(t.rows[0].cells, ["COLOUR", "MEANING", "USE"]):
        c.text = text
    rows = [
        ("Navy", "Stable information", "Navigation, major totals, identity"),
        ("Green", "Healthy actual progress", "Confirmed contribution or completed target"),
        ("Amber", "Attention soon", "Near spending limit, due item, review warning"),
        ("Red", "Real problem", "Overspending, failed sync/payment, overdue required reserve"),
        ("Grey", "Neutral progress", "Target still filling, inactive, informational"),
    ]
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            c.text = value
    b.style_table(t)
    b.add_heading(doc, "Progressive disclosure", 2)
    for item in [
        "List view answers what happened.",
        "Detail view answers why, who, and how it affected the plan.",
        "Settings answer what should happen next time.",
        "Do not place all three layers in one row or card.",
    ]:
        b.add_bullet(doc, item)
    b.add_heading(doc, "Design synthesis", 2)
    b.add_text(doc, "Use Wise's calm hierarchy as the baseline. Add Revolut's immediacy only at moments of action. Add Monarch's density only after the user asks for analysis. Keep Wollie's warm illustration and couples identity outside high-stakes numeric surfaces.")
    b.add_callout(doc, "The test", "A user should be able to explain any number in one sentence, then open it and see the exact inputs.", b.PALE_GREEN)
    doc.add_page_break()


def roadmap(doc):
    b.add_heading(doc, "Recommended product decisions", 1, before=0)
    b.add_heading(doc, "Build now", 2)
    for item in [
        "Household Overview with traceable totals.",
        "Transactions with totals, categories, Needs review, and simple rules.",
        "Money plan tabs for Spending this month and Savings & reserves.",
        "Proportional income allocation with a calculation preview.",
        "Enforced visibility: Private, Shared details, Household total only.",
    ]:
        b.add_number(doc, item)
    b.add_heading(doc, "Build next", 2)
    for item in [
        "Recurring detection with confidence and correction.",
        "Partner Light for a less hands-on partner.",
        "Contribution views for fixed, equal, and proportional sharing.",
        "Goal destination matching and actual contribution detection.",
        "Audit trail and undo for automatic rules.",
    ]:
        b.add_number(doc, item)
    b.add_heading(doc, "Avoid until the core is trusted", 2)
    for item in [
        "Moving money or executing transfers.",
        "Cards, investments, or a broad marketplace.",
        "A dashboard made of many configurable widgets.",
        "Forecasts that cannot be reconciled to source transactions.",
        "Gamification that treats overspending as shame.",
    ]:
        b.add_bullet(doc, item)
    b.add_callout(doc, "North star", "Both partners can answer: What can we spend? What did we save? What still needs attention?", b.PALE_ORANGE)
    doc.add_page_break()


def sources(doc):
    b.add_heading(doc, "Sources and evidence notes", 1, before=0)
    b.add_text(doc, "Official sources were accessed on 1 September 2026. Product availability, pricing, and regional rules may change. Annual-report screenshots are included as public product evidence. Reddit links are qualitative anecdotes, not prevalence estimates.", size=9.5, color=b.MUTED, italic=True)
    groups = [
        ("Wise official", [
            ("Spend with others help", "https://wise.com/help/articles/4nyqakqieuWseODSZAUCzi/how-to-use-spend-with-others"),
            ("Spend with others", "https://wise.com/p/spend-with-others"),
            ("Jars", "https://wise.com/help/articles/2978074/what-are-jars-and-how-can-i-keep-money-in-them"),
            ("Scheduled transfers", "https://wise.com/help/articles/2978063/what-are-scheduled-transfers"),
            ("Pricing", "https://wise.com/gb/pricing/"),
        ]),
        ("Revolut official", [
            ("Joint account setup", "https://help.revolut.com/en-BE/help/profile-and-plan/joint-accounts/how-to-set-up-a-joint-account-with-revolut/"),
            ("Income Sorter", "https://help.revolut.com/en-BE/help/accounts/budget-and-analytics/using-the-income-sorter-feature/"),
            ("Spending and income analytics", "https://help.revolut.com/en-NL/help/accounts/budget-and-analytics/how-can-i-see-my-spending-and-income-analytics/"),
            ("Custom categories", "https://help.revolut.com/en-IE/help/accounts/budget-and-analytics/how-to-create-and-manage-custom-categories-for-transactions/"),
            ("Belgian pricing", "https://www.revolut.com/en-BE/our-pricing-plans/"),
            ("FY2025 annual report", "https://assets.revolut.com/pdf/annualreport2025.pdf"),
        ]),
        ("Monarch official", [
            ("For couples", "https://www.monarch.com/for-couples"),
            ("Shared Views", "https://help.monarch.com/hc/en-us/articles/42228648365076-Shared-Views-in-Monarch"),
            ("Monarch for Couples", "https://help.monarch.com/hc/en-us/articles/20926382202004-Monarch-for-Couples"),
            ("Understanding your budget", "https://help.monarch.com/hc/en-us/articles/360048883631-Understanding-Your-Budget-in-Monarch"),
            ("Transaction rules", "https://help.monarchmoney.com/hc/en-us/articles/360048393372-Transaction-rules"),
            ("Reviewing transactions", "https://help.monarch.com/hc/en-us/articles/5528707082516-Reviewing-transactions"),
            ("Recurring expenses", "https://help.monarch.com/hc/en-us/articles/4890751141908-Tracking-Recurring-Expenses-and-Bills"),
            ("Reports", "https://help.monarch.com/hc/en-us/articles/21846787088916-Using-Reports"),
            ("Product updates", "https://www.monarch.com/whats-new"),
        ]),
    ]
    for heading, items in groups:
        b.add_heading(doc, heading, 2)
        for label, url in items:
            p = doc.add_paragraph(style="List Bullet")
            b.add_hyperlink(p, label, url)
            b.set_paragraph_spacing(p, after=1)
    doc.add_page_break()

    b.add_heading(doc, "Anecdotal user-discussion index", 1, before=0)
    links = [
        ("Wise and Revolut comparison", "https://www.reddit.com/r/ExpatFinance/comments/1rrl2iu/anyone_using_both_wise_and_revolut_accounts/"),
        ("Revolut joint-account setup confusion", "https://www.reddit.com/r/Revolut/comments/1ve5lvm/gb_joint_accounts/"),
        ("Revolut joint-account creation errors", "https://www.reddit.com/r/Revolut/comments/1dde3g2/i_cant_make_a_joint_account/"),
        ("Revolut analytics history complaint", "https://www.reddit.com/r/Revolut/comments/1n3k1wl/3_years_of_analytics_gone_after_the_last_update/"),
        ("Revolut Pockets analytics complaint", "https://www.reddit.com/r/Revolut/comments/1aw4hr5"),
        ("Monarch account-sync complaint", "https://www.reddit.com/r/MonarchMoney/comments/1vujxdx/account_syncing_is_the_product/"),
        ("Monarch duplicate transactions and trust", "https://www.reddit.com/r/MonarchMoney/comments/1b5mkwj/sync_issues_and_duplicate_transactions_undermine_trust_what_are_we_paying_for/"),
        ("Monarch private budgets and transactions request", "https://www.reddit.com/r/MonarchMoney/comments/1v4m9rv/private_budgets_and_transactions/"),
        ("Monarch goals criticism", "https://www.reddit.com/r/MonarchMoney/comments/1qgte8x/the_new_goals_system_is_a_massive_downgrade_and_i/"),
        ("Monarch budget criticism", "https://www.reddit.com/r/MonarchMoney/comments/1tdirbz/monarch_budget_30_is_the_worst_thing_ive_ever/"),
    ]
    for label, url in links:
        p = doc.add_paragraph(style="List Bullet")
        b.add_hyperlink(p, label, url)
        b.set_paragraph_spacing(p, after=2)
    b.add_heading(doc, "Research limits", 2)
    b.add_text(doc, "This is desk research, not a regulated account-opening or long-term sync test. Public screenshots may depict selected marketing states. Public discussions may be incomplete or unrepresentative. Validate the final Wollie flows with real couples using joint, separate, and mixed accounts.", after=0)


def build():
    doc = b.setup_document()
    doc.core_properties.title = "Wise, Revolut, and Monarch: UX, UI, and product logic deep dive"
    doc.core_properties.subject = "Competitive analysis and product recommendations for Wollie"
    doc.core_properties.author = "Wollie"
    # Running label, replacing inherited project text.
    header_table = doc.sections[0].header.tables[0]
    header_table.cell(0, 0).paragraphs[0].clear()
    r = header_table.cell(0, 0).paragraphs[0].add_run("WOLLIE  /  WISE · REVOLUT · MONARCH")
    b.set_run(r, 8, color=b.NAVY, bold=True)
    title(doc, "Wise, Revolut, and Monarch", "UX, UI, product logic, strengths, weaknesses, and decisions for Wollie")
    page_intro(doc)
    page_position(doc)
    page_product(
        doc,
        "Wise: clarity closest to the ledger",
        "The strongest reference for simple balances, transaction status, and confidence around money movement.",
        SHOTS / "wise-annual-report-page-024.png",
        "Wise's FY2025 annual report presents an Account screen built around balances and recent transactions.",
        [
            "The home model begins with accounts and currencies, then recent activity and actions.",
            "Jars separate actual money from the spendable balance.",
            "Spend with others delegates spending from an owner-controlled pool.",
            "Public materials emphasize operational money management rather than a household budget.",
        ],
        ["Clear hierarchy", "Strong status and fee trust cues", "Concrete real-money Jars", "Low transaction-list noise"],
        ["Not a true joint planning model", "Owner/member asymmetry", "Planning tied to money inside Wise", "No full cross-bank household view"],
    )
    wise_second(doc)
    page_product(
        doc,
        "Revolut: immediacy and breadth",
        "The strongest reference for actual-money automation, real joint accounts, and instant financial controls.",
        SHOTS / "revolut-annual-report-page-024.png",
        "Revolut's FY2025 annual report reflects a broad, high-energy super-app rather than a single planning workflow.",
        [
            "Personal and joint accounts are separate real ledger containers.",
            "Income Sorter can distribute actual incoming money to several destinations.",
            "Pockets, Savings, Investments, cards, and analytics create powerful but distributed states.",
            "Custom categories and multiple analysis dimensions support exploration.",
        ],
        ["True joint account", "Fast feedback and control", "Useful income automation", "Engaging, personalized UI"],
        ["Feature density", "Regional eligibility rules", "Joint analytics gaps", "Household picture can fragment"],
    )
    revolut_second(doc)
    monarch_first(doc)
    monarch_shared(doc)
    monarch_review(doc)
    monarch_budget(doc)
    journeys(doc)
    information_architecture(doc)
    visual_principles(doc)
    roadmap(doc)
    sources(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
