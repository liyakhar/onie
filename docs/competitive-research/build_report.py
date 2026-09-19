from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "wollie-couples-finance-research.docx"
SHOTS = ROOT / "screenshots"

# Document preset: standard_business_brief (Letter, 1in margins, Calibri 11pt).
# Cover pattern: editorial_cover, with Wollie's navy and orange used only as a named brand overlay.
NAVY = "082C4C"
ORANGE = "ED6A3A"
CREAM = "FFF8EF"
INK = "17212B"
MUTED = "667085"
BLUE = "2E74B5"
GREEN = "3A7D44"
RULE = "D9E2F3"
PALE_BLUE = "EEF5FC"
PALE_ORANGE = "FFF0E9"
PALE_GREEN = "EFF7F0"
PAGE_WIDTH_DXA = 9360  # 6.5in content width in standard_business_brief.


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, attrs in kwargs.items():
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in attrs.items():
            element.set(qn("w:{}".format(key)), str(value))


def set_table_geometry(table, widths, indent=120):
    """Explicit fixed table geometry in twips, including tblGrid and every cell."""
    assert sum(widths) == PAGE_WIDTH_DXA, (sum(widths), PAGE_WIDTH_DXA)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run(run, size=11, color=INK, bold=False, italic=False, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run(run, size=8, color=MUTED)


def add_text(doc, text, style=None, size=None, color=INK, bold=False, italic=False, align=None, before=0, after=6):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run(run, size=size or 11, color=color, bold=bold, italic=italic)
    set_paragraph_spacing(paragraph, before=before, after=after)
    return paragraph


def add_heading(doc, text, level=1, before=None, after=None):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run(run, size={1: 16, 2: 13, 3: 11}.get(level, 11), color=BLUE, bold=True)
    set_paragraph_spacing(paragraph, before=12 if before is None else before, after=4 if after is None else after)
    return paragraph


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run(run, 10.5)
    set_paragraph_spacing(paragraph, after=2)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_run(run, 10.5)
    set_paragraph_spacing(paragraph, after=2)
    return paragraph


def add_callout(doc, title, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": "6", "color": BLUE}, bottom={"val": "single", "sz": "6", "color": BLUE})
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    run = p.add_run(title + "  ")
    set_run(run, 10.5, color=NAVY, bold=True)
    run = p.add_run(text)
    set_run(run, 10.5, color=INK)
    set_paragraph_spacing(p, after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def style_table(table, header_fill=NAVY):
    for row_index, row in enumerate(table.rows):
        set_row_cant_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run(run, 9.5, color="FFFFFF", bold=True)
                    set_paragraph_spacing(paragraph, after=0)
            else:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run(run, 9, color=INK)
                    set_paragraph_spacing(paragraph, after=1)
    set_repeat_table_header(table.rows[0])


def add_picture_with_caption(doc, image_path, caption, width=6.05):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run(run, 8.5, color=MUTED, italic=True)
    set_paragraph_spacing(cap, after=8)


def add_rule(doc):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, RULE)
    cell.height = Inches(0.015)
    for paragraph in cell.paragraphs:
        set_paragraph_spacing(paragraph, after=0)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size in ((1, 16), (2, 13), (3, 11)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(4)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(MUTED)

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(2)

    core = doc.core_properties
    core.title = "Couples finance apps: product, UX and market research"
    core.subject = "Competitive research and product recommendations for Wollie"
    core.author = "Wollie research"
    core.keywords = "Wollie, couples, money planning, budgeting, UX, competitive research"

    # Header and footer for every page after the cover.
    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(header_table, [6000, 3360])
    left, right = header_table.rows[0].cells
    p = left.paragraphs[0]
    run = p.add_run("WOLLIE  /  COUPLES FINANCE RESEARCH")
    set_run(run, 8, color=NAVY, bold=True)
    set_paragraph_spacing(p, after=0)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("1 September 2026")
    set_run(run, 8, color=MUTED)
    set_paragraph_spacing(p, after=0)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(footer_table, [7200, 2160])
    p = footer_table.cell(0, 0).paragraphs[0]
    run = p.add_run("Wollie product research  •  private working document")
    set_run(run, 8, color=MUTED)
    set_paragraph_spacing(p, after=0)
    p = footer_table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(p)
    set_paragraph_spacing(p, after=0)
    return doc


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("WOLLIE  /  PRODUCT RESEARCH")
    set_run(run, 10, color=ORANGE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Couples finance apps:\nproduct, UX and market research")
    set_run(run, 28, color=NAVY, bold=True)
    p.paragraph_format.line_spacing = 1.03

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    run = p.add_run("What people use today, where it breaks, and the opportunity for Wollie")
    set_run(run, 14, color=MUTED)

    cover_table = doc.add_table(rows=1, cols=1)
    set_table_geometry(cover_table, [PAGE_WIDTH_DXA])
    cell = cover_table.cell(0, 0)
    set_cell_shading(cell, CREAM)
    set_cell_border(cell, top={"val": "single", "sz": "12", "color": ORANGE}, bottom={"val": "single", "sz": "12", "color": ORANGE})
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Scope: European couples using joint, separate, or mixed accounts\n1 September 2026")
    set_run(run, 11, color=NAVY, bold=True)
    set_paragraph_spacing(p, after=0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Research method")
    set_run(run, 10, color=ORANGE, bold=True)
    set_paragraph_spacing(p, after=3)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Official product documentation, public user discussions, and three safe zero-data trial accounts. No bank account, identity verification, money movement or paid subscription was used.")
    set_run(run, 10.5, color=MUTED)
    set_paragraph_spacing(p, after=0)

    doc.add_page_break()


def add_executive_summary(doc):
    add_heading(doc, "Decision in one page", 1, before=0)
    add_text(doc, "Wollie should be a couples-first money planning layer over the accounts people already use. It should work for one joint account, separate personal accounts, and a mixture of both.", size=12, color=NAVY, before=0, after=10)
    add_callout(doc, "Recommendation", "Do not position Wollie as another bank or a generic expense tracker. Own the space between a couple’s existing accounts and the shared decisions they need to make.", PALE_ORANGE)

    add_heading(doc, "The opportunity", 2)
    add_text(doc, "Across the market, couples repeatedly trade simplicity for control. A bank can hold shared money. A personal finance app can budget it. Neither reliably gives a couple one plan that respects personal boundaries, calculates contributions and stays useful when only one person wants to operate it.")

    cards = doc.add_table(rows=1, cols=3)
    set_table_geometry(cards, [3120, 3120, 3120])
    labels = [
        ("1", "Privacy by design", "Private, shared and household-total-only views. Never force full transaction visibility."),
        ("2", "Contribution rules", "Equal, fixed and income-proportional funding for shared spending and goals."),
        ("3", "Partner Light", "A calm safe-to-spend and approval view for the less hands-on partner."),
    ]
    for cell, (number, title, text) in zip(cards.rows[0].cells, labels):
        set_cell_shading(cell, PALE_BLUE)
        set_cell_border(cell, top={"val": "single", "sz": "6", "color": BLUE}, bottom={"val": "single", "sz": "6", "color": BLUE})
        p = cell.paragraphs[0]
        run = p.add_run(number + "\n")
        set_run(run, 16, color=ORANGE, bold=True)
        run = p.add_run(title + "\n")
        set_run(run, 10, color=NAVY, bold=True)
        run = p.add_run(text)
        set_run(run, 9, color=INK)
        set_paragraph_spacing(p, after=0)

    add_heading(doc, "A critical UI distinction", 2)
    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [4680, 4680])
    cells = table.rows[0].cells
    cells[0].text = "EVERYDAY SPENDING"
    cells[1].text = "FUTURE RESERVES"
    rows = [
        ("Food, travel, rent and discretionary categories.", "Savings, tax, pension, investment and future purchases."),
        ("A purchase reduces what is left to spend this month.", "Funding builds progress toward a target. It is not spending."),
    ]
    for r, values in zip(table.rows[1:], rows):
        for cell, value in zip(r.cells, values):
            cell.text = value
    style_table(table)
    add_text(doc, "Use red only for genuine overspending or an actionable overdue shortfall. A partially funded reserve is neutral progress, not a warning.", size=10.5, color=MUTED, italic=True, after=0)
    doc.add_page_break()


def add_landscape(doc):
    add_heading(doc, "What couples use today", 1, before=0)
    add_text(doc, "Direct planning products and adjacent banking apps solve different pieces of the same household problem. The comparison below focuses on their household model, not every individual feature.", color=MUTED, after=8)
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [1650, 2200, 2550, 2960])
    for cell, text in zip(table.rows[0].cells, ["PRODUCT", "BEST AT", "HOUSEHOLD MODEL", "KEY LESSON FOR WOLLIE"]):
        cell.text = text
    data = [
        ("YNAB", "Deep intentional budgeting.", "One subscription can share with up to five people. Joint, separate and mixed setups are possible.", "Offer one simpler default household plan and granular visibility."),
        ("Monarch", "One household dashboard across both partners’ accounts.", "Separate logins and shared data under one subscription.", "Never make account and transaction visibility all-or-nothing."),
        ("Wallet", "Clear personal dashboard and permissions.", "Premium owner creates a group; members switch between personal and group contexts.", "Do not force a context switch for ‘my’ versus ‘our’ money."),
        ("Spendee", "Simple wallet and category model.", "Shared wallet owned by one person; bank wallets cannot be shared.", "A shared plan must interpret both bank-linked and manual money."),
        ("Revolut", "Joint account behaviour, cards and Pockets.", "A true joint account for eligible customers.", "Meet the familiar bank mental model without becoming a bank."),
        ("Wise", "Cross-border money and controlled shared spending.", "Spend-with-others group model in supported markets.", "Be crystal-clear about ownership and what Wollie never does."),
    ]
    for item in data:
        row = table.add_row()
        for cell, text in zip(row.cells, item):
            cell.text = text
    style_table(table)
    add_callout(doc, "Market reading", "Revolut and Wise are important behaviour-setters, but they are not the same product category. Wollie should be the cross-bank planning layer that lets a couple understand money held elsewhere.", PALE_GREEN)

    add_heading(doc, "The main gaps", 2)
    add_bullet(doc, "A shared plan usually means broad visibility, multiple budgets or a separate shared-wallet context.")
    add_bullet(doc, "Contribution logic is often manual: who paid, who owes, and how income should fund the plan.")
    add_bullet(doc, "Savings, tax and pension are frequently mixed into a monthly budget as though they were spending.")
    add_bullet(doc, "The less-engaged partner is rarely given a purpose-built light experience.")
    doc.add_page_break()


def add_gallery(doc):
    add_heading(doc, "Trial-account UX gallery", 1, before=0)
    add_text(doc, "The following were captured in safe zero-data research accounts on 1 September 2026. No bank account was connected, no identity verification was submitted, and no money was moved. The gallery covers representative main screens, not every workflow.", color=MUTED, after=8)

    add_heading(doc, "YNAB: collaborative budget onboarding", 2)
    add_text(doc, "The onboarding-first experience introduces planning behaviour before showing a dense budget. It is friendly and intentional, although couples may still need to decide between multiple budget structures and manually manage personal contributions. [YNAB’s couples guide] supports joint, separate and mixed setups.", after=4)
    last = doc.paragraphs[-1]
    # Replace bracketed source text with a proper link in a short separate sentence for reliable DOCX output.
    last.clear()
    r = last.add_run("The onboarding-first experience introduces planning behaviour before showing a dense budget. It is friendly and intentional, although couples may still need to decide between multiple budget structures and manually manage personal contributions. ")
    set_run(r, 11)
    add_hyperlink(last, "YNAB’s couples guide", "https://www.ynab.com/guide/budgeting-as-a-couple")
    r = last.add_run(" supports joint, separate and mixed setups.")
    set_run(r, 11)
    set_paragraph_spacing(last, after=4)
    add_picture_with_caption(doc, SHOTS / "ynab-onboarding-cropped.png", "YNAB onboarding. Research account, zero data. Source: direct product observation, 1 September 2026.", width=6.05)
    doc.add_page_break()

    add_heading(doc, "Wallet by BudgetBakers: calm personal dashboard", 2)
    add_text(doc, "The navigation is easy to scan: Dashboard, Accounts, Records, Analytics and Imports. The familiar personal-money structure is a strength. Its group sharing model is more limiting: the official documentation says it is unavailable on the web, uses separate group context and initially gives new members access to all accounts.")
    add_picture_with_caption(doc, SHOTS / "wallet-dashboard.png", "Wallet dashboard. Research account, zero data. Source: direct product observation, 1 September 2026.", width=6.05)
    doc.add_page_break()

    add_heading(doc, "Wallet: account list", 2, before=0)
    add_text(doc, "The account list reinforces a simple tracker mental model. Wollie should take the calm clarity, while avoiding a separate ‘group wallet’ universe.")
    add_picture_with_caption(doc, SHOTS / "wallet-accounts.png", "Wallet accounts. Research account, zero data. Source: direct product observation, 1 September 2026.", width=6.05)
    doc.add_page_break()

    add_heading(doc, "Spendee: simple shared-wallet idea, early paywall", 2)
    add_text(doc, "Spendee makes wallets and categories approachable. In the trial account, adding a second wallet immediately reached a Premium boundary. The shared-wallet model is less complete than a shared household plan, and its help centre states that bank wallets cannot be shared.")
    add_picture_with_caption(doc, SHOTS / "spendee-wallet-paywall-sanitized.png", "Spendee Premium boundary. Profile label is intentionally redacted. Research account, zero data. Source: direct product observation, 1 September 2026.", width=5.85)
    doc.add_page_break()


def add_user_signals(doc):
    add_heading(doc, "What users repeatedly struggle with", 1, before=0)
    add_text(doc, "These are qualitative signals from public Reddit discussions. They are not survey results, but the themes recur across products and countries.", color=MUTED, after=8)

    signals = [
        ("Shared costs, personal autonomy", "Couples want to plan rent, groceries and trips together while retaining personal spending and surprise gifts. Full mutual visibility is often too blunt.", "Make visibility deliberate at account, transaction and category level: Private, Shared with partner, or Household total only."),
        ("One person runs the plan", "Many households have a ‘money person’ and a partner who does not want daily categorisation.", "Create a Partner Light role with safe-to-spend, upcoming bills, approvals and a weekly digest."),
        ("Separate cards create manual maths", "Shared cards, personal cards and paybacks create awkward contribution calculations.", "Support equal, fixed and proportional-to-income rules, then show expected and actual contributions."),
        ("People do not want to switch worlds", "Personal records and group records often become separate product contexts.", "Keep one household plan. Let privacy settings change the view rather than moving the user elsewhere."),
        ("The plan goes stale", "Large changes need a conversation, and unrecorded transfers make the plan feel arbitrary.", "Show simple proposals, confirmations and a lightweight decision history for material changes."),
    ]
    for number, (title, problem, implication) in enumerate(signals, start=1):
        table = doc.add_table(rows=1, cols=2)
        set_table_geometry(table, [1650, 7710])
        left, right = table.rows[0].cells
        set_cell_shading(left, PALE_ORANGE)
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(number))
        set_run(r, 18, color=ORANGE, bold=True)
        set_paragraph_spacing(p, after=0)
        p = right.paragraphs[0]
        r = p.add_run(title + "\n")
        set_run(r, 10.5, color=NAVY, bold=True)
        r = p.add_run(problem + "\n")
        set_run(r, 9.5, color=INK)
        r = p.add_run("Wollie implication: ")
        set_run(r, 9.5, color=GREEN, bold=True)
        r = p.add_run(implication)
        set_run(r, 9.5, color=INK)
        set_paragraph_spacing(p, after=0)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_heading(doc, "Examples of the public discussion", 2)
    p = doc.add_paragraph()
    add_hyperlink(p, "YNAB privacy and gifts", "https://www.reddit.com/r/ynab/comments/z2tt2y")
    p.add_run("  •  ")
    add_hyperlink(p, "YNAB shared-account setup", "https://www.reddit.com/r/ynab/comments/1tjdej2/how_to_budget_a_shared_bank_account/")
    p.add_run("  •  ")
    add_hyperlink(p, "Monarch automatic partner budget splits", "https://www.reddit.com/r/MonarchMoney/comments/160b0lw")
    p.add_run("  •  ")
    add_hyperlink(p, "Wallet joint and personal accounts", "https://www.reddit.com/r/BudgetBakers/comments/18wjv9b")
    for run in p.runs:
        set_run(run, 9, color=BLUE)
    set_paragraph_spacing(p, after=0)
    doc.add_page_break()


def add_ux_recommendations(doc):
    add_heading(doc, "UX principles for Wollie", 1, before=0)
    add_heading(doc, "1. Start with intent, not account architecture", 2)
    add_text(doc, "Ask: “Which money do you want to plan together?” Then let a couple include joint accounts, selected personal accounts, or both. Do not force a ‘joint versus separate’ decision before they see the model.")

    add_heading(doc, "2. One shared plan, layered views", 2)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120])
    for cell, (title, text) in zip(table.rows[0].cells, [
        ("MY VIEW", "My accounts and my private categories."),
        ("SHARED VIEW", "Money and categories both partners chose to share."),
        ("HOUSEHOLD VIEW", "Combined plan and totals, respecting privacy settings."),
    ]):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        r = p.add_run(title + "\n")
        set_run(r, 10, color=NAVY, bold=True)
        r = p.add_run(text)
        set_run(r, 9.5)
        set_paragraph_spacing(p, after=0)

    add_heading(doc, "3. Separate funding from spending", 2)
    add_text(doc, "At the top of Money plan, use two clear tabs or sections:")
    add_bullet(doc, "Spend this month: Food, travel, rent and discretionary categories. Example: ‘€800 left to spend.’")
    add_bullet(doc, "Save and prepare: Tax, pension, savings and future goals. Example: ‘€2,250 set aside’ and ‘45% of €5,000 target.’")
    add_text(doc, "A card purchase should update a spending category. A funded pension or tax reserve should build progress, not appear in ‘Spent this month.’", color=GREEN, bold=True)

    add_heading(doc, "4. Make contribution rules visible", 2)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120])
    for cell, (title, example) in zip(table.rows[0].cells, [
        ("EQUAL", "Food: funded 50/50."),
        ("FIXED", "Holiday: Liya €200, Alex €200 per month."),
        ("PROPORTIONAL", "Tax: funded in proportion to income."),
    ]):
        set_cell_shading(cell, PALE_GREEN)
        p = cell.paragraphs[0]
        r = p.add_run(title + "\n")
        set_run(r, 10, color=GREEN, bold=True)
        r = p.add_run(example)
        set_run(r, 9.5)
        set_paragraph_spacing(p, after=0)

    add_heading(doc, "5. Make card spending feel automatic, not mysterious", 2)
    add_text(doc, "Use one short explanation near transactions: “When a card purchase arrives, Wollie categorises it and updates that category’s balance.” Show the category chip, the transaction amount and the new remaining balance. Do not repeat the full envelope status below every transaction.")
    doc.add_page_break()


def add_priorities_and_flow(doc):
    add_heading(doc, "Recommended 90-day product priorities", 1, before=0)
    priorities = [
        ("Household scope and permissions", "Define Private, Shared and Household total only visibility. Make invite, accept, leave and revoke flows safe and simple."),
        ("Contribution and funding rules", "Support fixed, equal and proportional-to-income funding per spending category or goal. Keep a clear audit trail."),
        ("Spend versus reserve model", "Build distinct data and UI states for monthly spending and future reserves. Never label a pension allocation as spent."),
        ("Partner Light", "Create a quick, mobile-friendly view for safe-to-spend, bills, approvals and a weekly household summary."),
        ("Monthly planning conversation", "Add a low-friction check-in for income received, funding progress, proposed changes and decisions."),
    ]
    for number, (title, text) in enumerate(priorities, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{number}. {title}\n")
        set_run(r, 11, color=NAVY, bold=True)
        r = p.add_run(text)
        set_run(r, 10.5)
        set_paragraph_spacing(p, after=6)

    add_heading(doc, "Proposed permission model", 2)
    table = doc.add_table(rows=1, cols=5)
    set_table_geometry(table, [1750, 1700, 2000, 1900, 2010])
    headers = ["LEVEL", "BALANCE", "MERCHANT", "CATEGORY TOTAL", "EDIT PLAN"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    data = [
        ("Private", "Only owner", "Only owner", "Optional roll-up", "Owner only"),
        ("Shared", "Both partners", "Both partners", "Both partners", "By household role"),
        ("Household total only", "Included in total", "Hidden", "Shared impact only", "Owner controls source"),
    ]
    for values in data:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    style_table(table)
    add_text(doc, "The household-total-only option is strategically important. It lets a partner contribute to a shared plan without exposing every merchant.", size=10.5, color=GREEN, bold=True)

    add_heading(doc, "Invite flow", 2)
    for text in [
        "Invite by email.",
        "Recipient sees: “Sign in or create an account to accept.”",
        "If the address already has a Wollie account, accept without changing its login.",
        "Before connecting an account, choose: Private, Shared, or Household total only.",
        "Confirm what each partner can see and change.",
    ]:
        add_number(doc, text)
    doc.add_page_break()


def add_language_and_trust(doc):
    add_heading(doc, "Language, trust and competitive boundaries", 1, before=0)
    add_heading(doc, "Recommended customer language", 2)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [3350, 6010])
    table.rows[0].cells[0].text = "USE"
    table.rows[0].cells[1].text = "WHY"
    rows = [
        ("Money plan", "Broader and calmer than ‘budget’; includes spending, savings and goals."),
        ("Spend this month", "Plain label for day-to-day categories."),
        ("Save and prepare", "Makes tax, pension, savings and future goals feel different from spending."),
        ("Set aside / funded", "Correct language for reserves and future goals."),
        ("Available to spend", "Use only when clearly tied to everyday categories."),
    ]
    for row_data in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, row_data):
            cell.text = value
    style_table(table)

    add_heading(doc, "Words to avoid", 2)
    add_bullet(doc, "Avoid calling a pension, investment or savings allocation ‘spent.’")
    add_bullet(doc, "Avoid generic ‘household’ labels when the user needs to know whether an account is private, shared or total-only.")
    add_bullet(doc, "Avoid promising bank-like capabilities. Wollie does not hold or move money.")

    add_heading(doc, "Wollie’s trust promise", 2)
    add_callout(doc, "Suggested copy", "Wollie reads the accounts you choose. It never moves your money. You control what is shared.", PALE_GREEN)

    add_heading(doc, "What the banks teach us", 2)
    p = doc.add_paragraph()
    r = p.add_run("Revolut is a powerful local expectation-setter: its 2025 annual report says one million customers created Joint Accounts that year, and lists joint histories, matching cards, joint budgeting tools, Pockets and savings. ")
    set_run(r, 11)
    add_hyperlink(p, "Read the annual report", "https://assets.revolut.com/pdf/annualreport2025.pdf")
    set_paragraph_spacing(p, after=6)
    p = doc.add_paragraph()
    r = p.add_run("Wise is an adjacent controlled-sharing model. In supported regions the group owner retains legal control of the funds and members receive selected spending capabilities. ")
    set_run(r, 11)
    add_hyperlink(p, "Read Wise’s Spend with others guide", "https://wise.com/help/articles/4nyqakqieuWseODSZAUCzi/how-to-use-spend-with-others")
    set_paragraph_spacing(p, after=6)
    add_text(doc, "Wollie should borrow the clarity of account ownership and card spending, while remaining a calm cross-bank planning layer rather than a financial institution.", color=NAVY, bold=True, after=0)
    doc.add_page_break()


def add_sources(doc):
    add_heading(doc, "Sources and limitations", 1, before=0)
    add_heading(doc, "Official product sources", 2)
    add_text(doc, "All official source pages were accessed on 1 September 2026.", size=9.5, color=MUTED, italic=True, after=3)
    sources = [
        ("YNAB pricing and trial", "https://www.ynab.com/pricing/"),
        ("YNAB subscription sharing", "https://www.ynab.com/features/subscription-sharing"),
        ("YNAB guide to budgeting as a couple", "https://www.ynab.com/guide/budgeting-as-a-couple"),
        ("Monarch for Couples and Households", "https://help.monarch.com/hc/en-us/articles/20926382202004-Monarch-for-Couples-and-Households"),
        ("Monarch pricing", "https://help.monarch.com/hc/en-us/articles/9136169422996-Pricing"),
        ("Wallet group sharing", "https://support.budgetbakers.com/hc/en-us/articles/7149394922002-Everything-about-Group-Sharing"),
        ("Wallet overview", "https://support.budgetbakers.com/hc/en-us/articles/12212428113810-What-is-the-Wallet-app"),
        ("Spendee shared wallets", "https://help.spendee.com/article/224-shared-wallets"),
        ("Spendee Premium", "https://help.spendee.com/article/202-what-is-spendee-premium"),
        ("Revolut 2025 annual report", "https://assets.revolut.com/pdf/annualreport2025.pdf"),
        ("Wise: Spend with others", "https://wise.com/help/articles/4nyqakqieuWseODSZAUCzi/how-to-use-spend-with-others"),
    ]
    for title, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, title, url)
        set_paragraph_spacing(p, after=2)

    add_heading(doc, "Public user-discussion signals", 2)
    signals = [
        ("YNAB: couples who track expenses together", "https://www.reddit.com/r/ynab/comments/1ko2yhr/couples_who_track_expenses_together_how_do_you/"),
        ("YNAB: shared bank account setup", "https://www.reddit.com/r/ynab/comments/1tjdej2/how_to_budget_a_shared_bank_account/"),
        ("YNAB: privacy and gifts", "https://www.reddit.com/r/ynab/comments/z2tt2y"),
        ("YNAB: collaboration when one partner runs the budget", "https://www.reddit.com/r/ynab/comments/1swe8xf/need_advice_for_couples_collab_in_ynab/"),
        ("Monarch: household member options", "https://www.reddit.com/r/MonarchMoney/comments/1tzf5ub/options_for_household_members/"),
        ("Monarch: automatic partner budget splits", "https://www.reddit.com/r/MonarchMoney/comments/160b0lw"),
        ("Wallet: joint and personal accounts", "https://www.reddit.com/r/BudgetBakers/comments/18wjv9b"),
    ]
    for title, url in signals:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, title, url)
        set_paragraph_spacing(p, after=2)

    add_heading(doc, "Method limitations", 2)
    add_text(doc, "This is desk research plus three safe, zero-data trial accounts. It does not test live bank connections, regulated account opening, identity verification, paid subscriptions, support quality, account closure or long-term transaction classification. Public Reddit posts are qualitative and may be dated, incomplete or unrepresentative.", after=0)


def build():
    doc = setup_document()
    add_cover(doc)
    add_executive_summary(doc)
    add_landscape(doc)
    add_gallery(doc)
    add_user_signals(doc)
    add_ux_recommendations(doc)
    add_priorities_and_flow(doc)
    add_language_and_trust(doc)
    add_sources(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
